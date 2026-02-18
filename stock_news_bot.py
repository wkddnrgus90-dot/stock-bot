"""
주식 뉴스 자동 수집 봇
- 한국 주요 경제지 RSS 10분마다 수집
- Gemini API로 AI 분석 (심리/관련종목/테마섹터/핵심요약)
- 엑셀 저장 + 텔레그램 전송
- 중복 기사 필터링 (seen_articles.json)
"""

import os
import json
import logging
import time
from datetime import datetime
from pathlib import Path

import feedparser
import schedule
import requests
import google.generativeai as genai
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────
# 환경 변수 로드
# ──────────────────────────────────────────────
load_dotenv()

GEMINI_API_KEY    = os.getenv("GEMINI_API_KEY")
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
TELEGRAM_CHAT_ID  = os.getenv("TELEGRAM_CHAT_ID")

# ──────────────────────────────────────────────
# 로깅 설정
# ──────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler("bot.log", encoding="utf-8"),
    ],
)
logger = logging.getLogger(__name__)

# ──────────────────────────────────────────────
# RSS 피드 소스
# ──────────────────────────────────────────────
RSS_SOURCES = {
    "한국경제": "https://www.hankyung.com/feed/economy",
    "매일경제": "https://www.mk.co.kr/rss/30000001/",
    "연합뉴스": "https://www.yna.co.kr/RSS/economy.xml",
    "머니투데이": "https://rss.mt.co.kr/mt/1/rss.xml",
    "이데일리": "https://rss.edaily.co.kr/edaily/economy_news.xml",
    "인베스팅닷컴": "https://kr.investing.com/rss/news.rss",
}

# ──────────────────────────────────────────────
# 파일 경로
# ──────────────────────────────────────────────
SEEN_ARTICLES_FILE = Path("seen_articles.json")
EXCEL_FILE         = Path(f"stock_news_{datetime.now().strftime('%Y%m')}.xlsx")

# 엑셀 컬럼 정의
COLUMNS = ["No.", "시간", "채널(출처)", "심리", "관련종목", "테마섹터", "핵심요약", "원본링크"]
COL_WIDTHS = [6, 18, 12, 8, 20, 20, 60, 50]


# ──────────────────────────────────────────────
# Gemini API 초기화
# ──────────────────────────────────────────────
def init_gemini() -> genai.GenerativeModel:
    """Gemini 모델 초기화."""
    if not GEMINI_API_KEY:
        raise ValueError(".env 파일에 GEMINI_API_KEY가 설정되지 않았습니다.")
    genai.configure(api_key=GEMINI_API_KEY)
    return genai.GenerativeModel("gemini-1.5-flash")


# ──────────────────────────────────────────────
# 중복 기사 관리
# ──────────────────────────────────────────────
def load_seen_articles() -> set:
    """이미 처리한 기사 URL 목록 로드."""
    if SEEN_ARTICLES_FILE.exists():
        try:
            with open(SEEN_ARTICLES_FILE, "r", encoding="utf-8") as f:
                return set(json.load(f))
        except (json.JSONDecodeError, IOError) as e:
            logger.warning(f"seen_articles.json 로드 실패, 새로 시작: {e}")
    return set()


def save_seen_articles(seen: set) -> None:
    """처리한 기사 URL 목록 저장."""
    with open(SEEN_ARTICLES_FILE, "w", encoding="utf-8") as f:
        json.dump(list(seen), f, ensure_ascii=False, indent=2)


# ──────────────────────────────────────────────
# RSS 수집
# ──────────────────────────────────────────────
def fetch_rss(source_name: str, url: str) -> list[dict]:
    """단일 RSS 피드에서 기사 목록 수집."""
    articles = []
    try:
        feed = feedparser.parse(url)
        for entry in feed.entries:
            title   = entry.get("title", "").strip()
            link    = entry.get("link", "").strip()
            summary = entry.get("summary", entry.get("description", "")).strip()

            # 발행 시각 파싱
            published = ""
            if hasattr(entry, "published_parsed") and entry.published_parsed:
                published = datetime(*entry.published_parsed[:6]).strftime("%Y-%m-%d %H:%M")
            else:
                published = datetime.now().strftime("%Y-%m-%d %H:%M")

            if title and link:
                articles.append({
                    "source":    source_name,
                    "title":     title,
                    "link":      link,
                    "summary":   summary,
                    "published": published,
                })
    except Exception as e:
        logger.error(f"[{source_name}] RSS 수집 오류: {e}")
    return articles


def fetch_all_rss(seen: set) -> list[dict]:
    """모든 RSS 소스에서 신규 기사 수집."""
    new_articles = []
    for name, url in RSS_SOURCES.items():
        articles = fetch_rss(name, url)
        for article in articles:
            if article["link"] not in seen:
                new_articles.append(article)
        logger.info(f"[{name}] {len(articles)}개 수집, "
                    f"신규: {sum(1 for a in articles if a['link'] not in seen)}개")
    return new_articles


# ──────────────────────────────────────────────
# Gemini AI 분석
# ──────────────────────────────────────────────
ANALYSIS_PROMPT_TEMPLATE = """
다음 한국 경제/주식 뉴스 기사를 분석해 주세요.

제목: {title}
내용: {summary}

아래 형식으로 정확히 답변해 주세요. 각 항목은 반드시 한 줄로 작성하세요.

심리: [긍정/부정/중립 중 하나]
관련종목: [관련 종목명 또는 종목코드, 없으면 "해당없음"]
테마섹터: [관련 테마 또는 섹터명, 없으면 "기타"]
핵심요약: [2~3문장으로 핵심 내용 요약]
""".strip()


def analyze_article(model: genai.GenerativeModel, article: dict) -> dict:
    """Gemini API로 기사 분석."""
    prompt = ANALYSIS_PROMPT_TEMPLATE.format(
        title=article["title"],
        summary=article["summary"] or article["title"],
    )
    try:
        response = model.generate_content(prompt)
        return parse_gemini_response(response.text)
    except Exception as e:
        logger.error(f"Gemini 분석 오류 ({article['title'][:30]}...): {e}")
        return {"심리": "분석실패", "관련종목": "-", "테마섹터": "-", "핵심요약": article["title"]}


def parse_gemini_response(text: str) -> dict:
    """Gemini 응답 텍스트 파싱."""
    result = {"심리": "중립", "관련종목": "해당없음", "테마섹터": "기타", "핵심요약": ""}
    for line in text.strip().splitlines():
        if ":" in line:
            key, _, value = line.partition(":")
            key   = key.strip()
            value = value.strip()
            if key in result:
                result[key] = value
    return result


# ──────────────────────────────────────────────
# 엑셀 저장
# ──────────────────────────────────────────────
HEADER_FILL  = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT  = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
BODY_FONT    = Font(name="맑은 고딕", size=9)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

SENTIMENT_COLORS = {
    "긍정": "C6EFCE",  # 연초록
    "부정": "FFC7CE",  # 연빨강
    "중립": "FFEB9C",  # 연노랑
}

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def get_or_create_workbook() -> tuple[Workbook, object]:
    """엑셀 파일 로드 또는 신규 생성."""
    if EXCEL_FILE.exists():
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "주식뉴스"
        _write_header(ws)
    return wb, ws


def _write_header(ws) -> None:
    """엑셀 헤더 행 작성."""
    ws.append(COLUMNS)
    for col_idx, (_, width) in enumerate(zip(COLUMNS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border    = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def save_to_excel(articles_with_analysis: list[dict]) -> int:
    """분석된 기사를 엑셀에 저장. 저장된 건수 반환."""
    if not articles_with_analysis:
        return 0

    wb, ws = get_or_create_workbook()
    saved_count = 0

    for item in articles_with_analysis:
        article  = item["article"]
        analysis = item["analysis"]

        no       = ws.max_row  # 헤더(1행) 포함이므로 현재 max_row = 이전 마지막행
        row_data = [
            no,
            article["published"],
            article["source"],
            analysis["심리"],
            analysis["관련종목"],
            analysis["테마섹터"],
            analysis["핵심요약"],
            article["link"],
        ]
        ws.append(row_data)
        row_idx = ws.max_row

        # 스타일 적용
        sentiment = analysis.get("심리", "중립")
        fill_color = SENTIMENT_COLORS.get(sentiment, "FFFFFF")
        row_fill   = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        for col_idx in range(1, len(COLUMNS) + 1):
            cell            = ws.cell(row=row_idx, column=col_idx)
            cell.font       = BODY_FONT
            cell.border     = THIN_BORDER
            cell.fill       = row_fill
            cell.alignment  = CENTER_ALIGN if col_idx in (1, 2, 3, 4) else LEFT_ALIGN

        ws.row_dimensions[row_idx].height = 40
        saved_count += 1

    wb.save(EXCEL_FILE)
    logger.info(f"엑셀 저장 완료: {saved_count}건 → {EXCEL_FILE}")
    return saved_count


# ──────────────────────────────────────────────
# 텔레그램 전송
# ──────────────────────────────────────────────
def send_telegram(message: str) -> bool:
    """텔레그램 메시지 전송."""
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        logger.warning("텔레그램 설정 미완료 (.env 확인 필요)")
        return False

    url     = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
    payload = {
        "chat_id":    TELEGRAM_CHAT_ID,
        "text":       message,
        "parse_mode": "HTML",
        "disable_web_page_preview": True,
    }
    try:
        resp = requests.post(url, json=payload, timeout=10)
        resp.raise_for_status()
        return True
    except requests.RequestException as e:
        logger.error(f"텔레그램 전송 오류: {e}")
        return False


def build_telegram_message(articles_with_analysis: list[dict]) -> str:
    """텔레그램 요약 메시지 생성."""
    now   = datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = [f"📰 <b>주식 뉴스 요약</b> ({now})\n총 {len(articles_with_analysis)}건\n"]

    SENTIMENT_EMOJI = {"긍정": "🟢", "부정": "🔴", "중립": "🟡"}

    for i, item in enumerate(articles_with_analysis[:10], start=1):  # 최대 10건
        article  = item["article"]
        analysis = item["analysis"]
        emoji    = SENTIMENT_EMOJI.get(analysis["심리"], "⚪")

        lines.append(
            f"{i}. {emoji} <b>[{article['source']}]</b> {article['title']}\n"
            f"   📌 {analysis['핵심요약']}\n"
            f"   🏷 {analysis['테마섹터']} | {analysis['관련종목']}\n"
            f"   🔗 <a href=\"{article['link']}\">기사 보기</a>\n"
        )

    if len(articles_with_analysis) > 10:
        lines.append(f"\n... 외 {len(articles_with_analysis) - 10}건 (엑셀 파일 확인)")

    return "\n".join(lines)


# ──────────────────────────────────────────────
# 메인 실행 루틴
# ──────────────────────────────────────────────
def run_bot(model: genai.GenerativeModel) -> None:
    """봇 메인 실행 함수 (매 10분마다 호출)."""
    logger.info("=" * 50)
    logger.info("뉴스 수집 시작")

    seen     = load_seen_articles()
    new_arts = fetch_all_rss(seen)

    if not new_arts:
        logger.info("신규 기사 없음")
        return

    logger.info(f"신규 기사 {len(new_arts)}건 분석 시작")
    results = []
    for idx, article in enumerate(new_arts, start=1):
        logger.info(f"  [{idx}/{len(new_arts)}] {article['source']}: {article['title'][:40]}...")
        analysis = analyze_article(model, article)
        results.append({"article": article, "analysis": analysis})
        seen.add(article["link"])
        time.sleep(1)  # API 호출 간격 (과부하 방지)

    # 엑셀 저장
    save_to_excel(results)

    # 텔레그램 전송
    msg = build_telegram_message(results)
    if send_telegram(msg):
        logger.info("텔레그램 전송 완료")

    # 처리된 기사 목록 저장
    save_seen_articles(seen)
    logger.info("사이클 완료\n")


# ──────────────────────────────────────────────
# 환경 변수 유효성 검사
# ──────────────────────────────────────────────
def validate_env() -> bool:
    """필수 환경 변수 확인."""
    missing = []
    if not GEMINI_API_KEY:
        missing.append("GEMINI_API_KEY")
    if not TELEGRAM_BOT_TOKEN:
        missing.append("TELEGRAM_BOT_TOKEN")
    if not TELEGRAM_CHAT_ID:
        missing.append("TELEGRAM_CHAT_ID")

    if missing:
        logger.error(f"누락된 환경 변수: {', '.join(missing)}")
        logger.error(".env 파일을 확인하거나 .env.example을 참고하세요.")
        return False
    return True


# ──────────────────────────────────────────────
# 엔트리포인트
# ──────────────────────────────────────────────
if __name__ == "__main__":
    if not validate_env():
        exit(1)

    try:
        model = init_gemini()
        logger.info("Gemini 모델 초기화 완료")
    except Exception as e:
        logger.error(f"Gemini 초기화 실패: {e}")
        exit(1)

    logger.info("주식 뉴스 봇 시작 (10분 간격 자동 수집)")

    # 최초 즉시 실행
    run_bot(model)

    # 10분 간격 스케줄 등록
    schedule.every(10).minutes.do(run_bot, model=model)

    while True:
        schedule.run_pending()
        time.sleep(30)
